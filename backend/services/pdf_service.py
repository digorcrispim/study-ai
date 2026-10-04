from io import BytesIO

from pypdf import PdfReader
from pdf2image import convert_from_bytes
import pytesseract

MINIMUM_TEXT_LENGTH = 50


def extract_text_from_pdf(pdf_bytes: bytes) -> str:
    """Extrai texto de um PDF de forma resiliente.

    Estratégia:
      1. Tenta extração nativa com pypdf (rápida, para PDFs com texto).
      2. Se o texto nativo for vazio ou curto demais (< MINIMUM_TEXT_LENGTH),
         faz fallback para OCR (pdf2image + pytesseract), útil para PDFs
         escaneados/imagem.
      3. Se nada funcionar, lança ValueError com mensagem clara.
    """
    # 1. Tentativa nativa (pypdf)
    try:
        reader = PdfReader(BytesIO(pdf_bytes))

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""
            text = text.strip()

            if text:
                pages.append(text)

        native_text = "\n\n".join(pages).strip()

        if len(native_text) > MINIMUM_TEXT_LENGTH:
            return native_text
    except Exception as exc:
        print(f"Extração nativa falhou: {exc}")

    # 2. Fallback OCR (pdf2image + pytesseract)
    try:
        images = convert_from_bytes(pdf_bytes)

        ocr_pages = []

        for image in images:
            ocr_text = pytesseract.image_to_string(image, lang="por")
            ocr_text = ocr_text.strip()

            if ocr_text:
                ocr_pages.append(ocr_text)

        ocr_result = "\n\n".join(ocr_pages).strip()

        if len(ocr_result) > MINIMUM_TEXT_LENGTH:
            return ocr_result
    except Exception as exc:
        raise ValueError(
            f"Falha na extração por OCR: {exc}. "
            "Verifique se o Tesseract-OCR está instalado no sistema."
        )

    raise ValueError(
        "Não foi possível extrair texto do PDF. O arquivo pode estar "
        "vazio, corrompido ou ser uma imagem não legível."
    )


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extrai texto de arquivo .txt (decodifica como UTF-8)."""
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        raise ValueError(
            "Não foi possível decodificar o arquivo .txt. "
            "Verifique se está em UTF-8."
        )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extrai texto de arquivo .docx usando python-docx."""
    import io
    from docx import Document

    try:
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])

        if len(text.strip()) < MINIMUM_TEXT_LENGTH:
            raise ValueError(
                "O arquivo .docx parece estar vazio ou com muito pouco texto."
            )

        return text.strip()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Falha ao extrair texto do .docx: {exc}")


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    """Extrai texto de planilha .xlsx (Excel) usando openpyxl."""
    import io
    from openpyxl import load_workbook

    try:
        wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
        text_parts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join(
                    str(cell) for cell in row if cell is not None
                )
                if row_text.strip():
                    text_parts.append(row_text)
        return "\n".join(text_parts).strip()
    except Exception as exc:
        raise ValueError(f"Falha ao extrair texto do .xlsx: {exc}")


def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extrai texto de apresentação .pptx (PowerPoint) usando python-pptx."""
    import io
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(file_bytes))
        text_parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text_parts.append(shape.text)
        return "\n".join(text_parts).strip()
    except Exception as exc:
        raise ValueError(f"Falha ao extrair texto do .pptx: {exc}")


def extract_text_from_odf(file_bytes: bytes, content_type: str) -> str:
    """Extrai texto de documentos LibreOffice (.odt, .ods, .odp) usando odfpy."""
    import io
    from odf.opendocument import load
    from odf.text import P
    from odf import teletype

    try:
        doc = load(io.BytesIO(file_bytes))
        text_parts = []
        for paragraph in doc.getElementsByType(P):
            # teletype.extractText devolve o texto do nó (str(p) retornaria
            # a representação XML do elemento, não o conteúdo).
            text = teletype.extractText(paragraph)
            if text and text.strip():
                text_parts.append(text)
        return "\n".join(text_parts).strip()
    except Exception as exc:
        raise ValueError(f"Falha ao extrair texto do documento ODF: {exc}")


def extract_text_from_document(
    file_bytes: bytes, content_type: str, filename: str = "document"
) -> str:
    """Despacha para a função de extração correta baseada no content_type."""
    odf_types = {
        "application/vnd.oasis.opendocument.text",
        "application/vnd.oasis.opendocument.spreadsheet",
        "application/vnd.oasis.opendocument.presentation",
    }

    if content_type == "application/pdf":
        return extract_text_from_pdf(file_bytes)
    elif content_type == "text/plain":
        return extract_text_from_txt(file_bytes)
    elif content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return extract_text_from_docx(file_bytes)
    elif content_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        return extract_text_from_xlsx(file_bytes)
    elif content_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ):
        return extract_text_from_pptx(file_bytes)
    elif content_type in odf_types:
        return extract_text_from_odf(file_bytes, content_type)
    elif content_type.startswith("video/"):
        from .video_service import extract_text_from_video
        return extract_text_from_video(file_bytes, filename)
    else:
        raise ValueError(f"Formato de arquivo não suportado: {content_type}")
